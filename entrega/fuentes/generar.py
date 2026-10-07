"""Pasa los manuales y el acta de El Patio al diseño de la propuesta.

Lee los .docx de entrega/ (el contenido sigue viviendo ahí), arma un HTML por
documento con la misma hoja de estilo de «Propuesta - Sistema El Patio.html»
y lo imprime a PDF con Edge sin ventana. Las páginas las reparte un pequeño
guion dentro del HTML, así cada documento puede crecer sin descuadrarse.

Uso:  python generar.py <carpeta de salida de los PDF>
"""
import html
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

AQUI = Path(__file__).resolve().parent
ENTREGA = AQUI.parent
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
LUGAR = "El Patio · Turbaco, Bolívar"

# El logo y la hoja de estilo salen de la propuesta, que es la referencia.
_propuesta = (AQUI / "Propuesta - Sistema El Patio.html").read_text(encoding="utf-8")
LOGO = re.search(r'<img src="(data:image/png;base64,[^"]+)"', _propuesta).group(1)
CSS_PROPUESTA = re.search(r"<style>(.*?)</style>", _propuesta, re.S).group(1)

CSS_EXTRA = """
  .pagina { display: flex; flex-direction: column; }
  .pagina main { flex: 1; }
  .bajada { color: var(--gris); font-size: 9.5pt; margin: 0 0 16px; }
  .paso { display: grid; grid-template-columns: 17mm 1fr; gap: 0 4mm; margin: 10px 0; break-inside: avoid; }
  .paso .n { color: var(--cobre); font-size: 8pt; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; padding-top: 2px; border-top: 2px solid var(--cobre-claro); }
  .paso .t { font-weight: 700; margin: 0 0 2px; }
  .paso p { margin: 0; }
  main > p { margin: 0 0 8px; }
  ul.acuerdos { margin: 4px 0 8px; }
  ul.acuerdos li { margin: 0 0 5px; }
  ul.acuerdos li::marker { color: var(--cobre-claro); }
  .ficha { margin: 0 0 14px; }
  .ficha td:first-child { width: 32%; font-weight: 700; color: var(--cobre); font-size: 8.5pt; letter-spacing: .04em; text-transform: uppercase; background: var(--crema-suave); }
  .renglon { border-bottom: 1px solid var(--gris); height: 9mm; }
  td.vacia { height: 9mm; }
  .portada img { height: 150px; }
  #flujo { display: none; }
"""

# Reparte los bloques de #flujo en páginas carta. Un título nunca queda solo al
# pie: si lo que le sigue no cabe, se va con él a la página siguiente. Las
# tablas y las listas largas se parten por filas, repitiendo el encabezado.
PAGINADOR = """
<script>
// Se espera a que cargue el logo: antes de eso la portada mide menos.
window.addEventListener('load', function () {
  const flujo = document.getElementById('flujo');
  const tpl = document.getElementById('plantilla').content;
  const paginas = [];
  function nueva() {
    const p = tpl.cloneNode(true).firstElementChild;
    if (paginas.length === 0) p.querySelector('.banda').remove();
    else p.querySelector('.portada').remove();
    document.body.appendChild(p);
    paginas.push(p);
    return p.querySelector('main');
  }
  function sobra(main) {
    const pie = main.parentElement.querySelector('.pie').getBoundingClientRect().top - 18;
    const ult = main.lastElementChild;
    return ult && ult.getBoundingClientRect().bottom > pie;
  }
  // Títulos, y el párrafo corto que presenta una lista o tabla, no se quedan solos al pie.
  const atado = el => el.classList.contains('numero') || /^H[1-3]$/.test(el.tagName) || el.classList.contains('rotulo')
    || (el.tagName === 'P' && el.textContent.length < 260);
  let main = nueva();
  for (const bloque of Array.from(flujo.children)) {
    main.appendChild(bloque);
    if (!sobra(main)) continue;
    const partible = bloque.tagName === 'TABLE' ? bloque.tBodies[0] : (bloque.tagName === 'UL' ? bloque : null);
    let resto = null;
    if (partible && partible.children.length > 1) {
      resto = bloque.cloneNode(true);
      const destino = resto.tagName === 'TABLE' ? resto.tBodies[0] : resto;
      destino.innerHTML = '';
      while (sobra(main) && partible.children.length > 1) destino.prepend(partible.lastElementChild);
      if (sobra(main) || partible.children.length < (bloque.tagName === 'TABLE' ? 3 : 2)) {   // no cabe ni el comienzo: va entera
        while (destino.firstElementChild) partible.appendChild(destino.firstElementChild);
        resto = null;
      }
    }
    if (resto) { main = nueva(); main.appendChild(resto); continue; }
    const llevar = [bloque];
    while (llevar.length < 4 && main.children.length > llevar.length && atado(main.children[main.children.length - llevar.length - 1]))
      llevar.unshift(main.children[main.children.length - llevar.length - 1]);
    if (llevar.length === main.children.length) continue;  // ocupa la página entera: se queda
    main = nueva();
    for (const el of llevar) main.appendChild(el);
  }
  flujo.remove();
  paginas.forEach((p, i) => {
    p.querySelector('.num').textContent = 'Página ' + (i + 1) + ' de ' + paginas.length;
    const m = p.querySelector('main'); if (m.firstElementChild) m.firstElementChild.style.marginTop = '6px';
  });
});
</script>
"""


def esc(t):
    return html.escape(t, quote=False)


def runs_html(p):
    """Texto del párrafo conservando las negritas."""
    partes = []
    for r in p.runs:
        t = esc(r.text)
        if not t:
            continue
        partes.append(f"<b>{t}</b>" if r.bold else t)
    return re.sub(r"</b><b>", "", "".join(partes)).strip()


def frase(t):
    """«PUEDE ENVIAR VARIAS VECES» → «Puede enviar varias veces»."""
    t = t.strip().lower()
    return t[:1].upper() + t[1:]


def bloques(path):
    d = docx.Document(path)
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, d)
            if p.text.strip():
                yield "p", p
        elif tag == "tbl":
            yield "t", Table(el, d)


def documento(path):
    """Convierte un manual (todos tienen la misma estructura) en HTML de bloques."""
    salida, rotulo, titulo = [], None, None
    pendiente_paso = None
    lista = []

    def cerrar_lista():
        if lista:
            salida.append('<ul class="acuerdos">' + "".join(f"<li>{x}</li>" for x in lista) + "</ul>")
            lista.clear()

    for tipo, b in bloques(path):
        if tipo == "p":
            texto = b.text.strip()
            negrita = bool(b.runs) and bool(b.runs[0].bold)
            if texto.startswith("R E S T A U R A N T E") or texto.startswith("Turbaco, Bolívar"):
                continue
            if b.style.name == "List Bullet":
                lista.append(runs_html(b))
                continue
            cerrar_lista()
            if rotulo is None and negrita:
                rotulo = texto
                continue
            if titulo is None and negrita:
                titulo = texto
                continue
            m = re.match(r"^Paso (\d+)\.\s+(.*)$", texto)
            if m and negrita:
                pendiente_paso = (m.group(1), esc(m.group(2)))
                continue
            if pendiente_paso:
                n, t = pendiente_paso
                salida.append(f'<div class="paso"><div class="n">Paso {n}</div><div><p class="t">{t}</p><p>{runs_html(b)}</p></div></div>')
                pendiente_paso = None
                continue
            m = re.match(r"^(\d+)\.\s{2,}(.*)$", texto)
            if m and negrita:
                salida.append(f'<p class="numero">{int(m.group(1)):02d}</p>')
                salida.append(f"<h2>{esc(m.group(2))}</h2>")
                continue
            if negrita and len(b.runs) == 1 or (negrita and texto == "".join(r.text for r in b.runs if r.bold).strip()):
                salida.append(f"<h3>{esc(texto)}</h3>")
                continue
            if salida and salida[-1].startswith('<div class="paso">'):  # segundo párrafo del mismo paso
                salida[-1] = salida[-1][:-len("</div></div>")] + f"<p style=\"margin-top:6px\">{runs_html(b)}</p></div></div>"
                continue
            salida.append(f"<p>{runs_html(b)}</p>")
        else:
            cerrar_lista()
            filas = b.rows
            if len(filas) == 1 and len(set(id(c._tc) for c in filas[0].cells)) == 1:
                ps = [p for p in filas[0].cells[0].paragraphs if p.text.strip()]
                if len(ps) == 1:  # la frase bajo el título del documento
                    salida.append(f'<p class="bajada">{esc(ps[0].text.strip())}</p>')
                else:             # recuadro destacado: título en mayúsculas + texto
                    cuerpo = " ".join(esc(p.text.strip()) for p in ps[1:])
                    salida.append(f'<div class="aviso"><b>{esc(frase(ps[0].text))}.</b> {cuerpo}</div>')
                continue
            cab = [esc(c.text.strip()) for c in filas[0].cells]
            anchos = {2: ["34%", None], 3: ["22%", "30%", None]}.get(len(cab), [None] * len(cab))
            th = "".join(f'<th style="width:{w}">{c}</th>' if w else f"<th>{c}</th>" for c, w in zip(cab, anchos))
            cuerpo = "".join(
                "<tr>" + "".join(
                    f'<td class="fuerte">{esc(c.text.strip())}</td>' if i == 0 else f"<td>{esc(c.text.strip())}</td>"
                    for i, c in enumerate(r.cells)) + "</tr>"
                for r in filas[1:])
            salida.append(f"<table><thead><tr>{th}</tr></thead><tbody>{cuerpo}</tbody></table>")
    cerrar_lista()
    rotulo = rotulo.capitalize() if rotulo.isupper() else rotulo
    cab = [f'<p class="rotulo">{esc(rotulo)}</p>', f"<h1>{esc(titulo)}</h1>"]
    return cab, salida, f"{rotulo} · {titulo}"


def acta():
    """El acta se arma aparte: el docx trae campos para llenar a mano y firmas,
    y en el PDF solo va lo que lee el dueño (sin renglones ni firmas)."""
    path = ENTREGA / "Acta de entrega - Sistema El Patio.docx"
    d = docx.Document(path)
    tablas = d.tables
    ficha = "".join(
        f"<tr><td>{esc(r.cells[0].text.strip())}</td><td>{'' if '___' in r.cells[1].text else esc(r.cells[1].text.strip())}</td></tr>"
        for r in tablas[0].rows if r.cells[0].text.strip() and "___" not in r.cells[1].text)
    cab = ['<p class="rotulo">Acta de entrega</p>', "<h1>Sistema de gestión y sitio web para El Patio</h1>"]
    out = [f'<table class="ficha"><tbody>{ficha}</tbody></table>']

    lista = []
    def cerrar():
        if lista:
            out.append('<ul class="acuerdos">' + "".join(f"<li>{x}</li>" for x in lista) + "</ul>")
            lista.clear()

    parrafos = [p for p in d.paragraphs if p.text.strip()][2:]
    for p in parrafos:
        t = p.text.strip()
        if p.style.name == "List Bullet":
            lista.append(runs_html(p)); continue
        cerrar()
        m = re.match(r"^(\d+)\.\s+(.*)$", t)
        if m and p.runs[0].bold:
            num = int(m.group(1))
            if num >= 6:   # Observaciones y Firmas: son para llenar a mano
                break
            out.append(f'<p class="numero">{num:02d}</p>')
            out.append(f"<h2>{esc(m.group(2))}</h2>")
            continue
        if t.startswith("____"):
            continue
        if "constancia de que recibió" in t:
            out.append('<p class="intro">Con este documento se deja constancia de que el restaurante recibió el sistema '
                       '«El Patio» instalado y funcionando, y de que se le explicó su uso. Desde la entrega empiezan a '
                       'contar los doce meses de soporte incluidos en la propuesta aprobada.</p>'); continue
        out.append(f"<p>{runs_html(p)}</p>")
        # Las tablas van justo después del párrafo que las presenta.
        if t.startswith("Estas cuentas quedan"):  # la de capacitación es solo para llenar
            tb = tablas[2]
            cabs = [esc(c.text.strip()) for c in tb.rows[0].cells][:2]  # sin la columna «Titular» en blanco
            anchos = ["38%", None]
            th = "".join(f'<th style="width:{w}">{c}</th>' if w else f"<th>{c}</th>" for c, w in zip(cabs, anchos))
            filas = ""
            for r in tb.rows[1:]:
                celdas = [c.text.strip() for c in r.cells][:2]
                filas += "<tr>" + "".join(
                    f'<td class="fuerte">{esc(c)}</td>' if i == 0 else
                    (f'<td class="vacia"></td>' if (not c or "___" in c) else f"<td>{esc(c)}</td>")
                    for i, c in enumerate(celdas)) + "</tr>"
            out.append(f"<table><thead><tr>{th}</tr></thead><tbody>{filas}</tbody></table>")
    cerrar()
    return cab, out, "Acta de entrega"


def pagina_html(titulo_doc, cab, cuerpo):
    return f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<title>{esc(titulo_doc)} · El Patio</title>
<style>{CSS_PROPUESTA}{CSS_EXTRA}</style>
</head>
<body>
<template id="plantilla"><section class="pagina">
  <header class="portada"><img src="{LOGO}" alt="El Patio"></header>
  <div class="banda">El Patio &nbsp;·&nbsp; {esc(titulo_doc)}</div>
  <main></main>
  <div class="pie"><span>{esc(titulo_doc)} · {LUGAR}</span><span class="num"></span></div>
</section></template>
<div id="flujo">
{chr(10).join(cab)}
{chr(10).join(cuerpo)}
</div>
{PAGINADOR}
</body>
</html>
"""


def imprimir(html_path, pdf_path, perfil):
    # Edge sin ventana: la ruta del PDF va sin espacios y se renombra después.
    tmp = perfil.parent / "salida.pdf"
    tmp.unlink(missing_ok=True)
    subprocess.run([EDGE, "--headless=new", "--disable-gpu", f"--user-data-dir={perfil}",
                    "--no-pdf-header-footer", "--virtual-time-budget=4000",
                    f"--print-to-pdf={tmp}", html_path.as_uri()], check=True, timeout=120,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    tmp.replace(pdf_path)


def main():
    destino = Path(sys.argv[1])
    perfil = Path(sys.argv[2]) if len(sys.argv) > 2 else AQUI / ".edge-perfil"
    trabajos = [(acta(), "Acta de entrega - Sistema El Patio", destino)]
    for f in sorted((ENTREGA / "manuales").glob("*.docx")):
        trabajos.append((documento(f), f.stem, destino / "Manuales"))
    for (cab, cuerpo, titulo), nombre, carpeta in trabajos:
        h = AQUI / f"{nombre}.html"
        h.write_text(pagina_html(titulo, cab, cuerpo), encoding="utf-8")
        imprimir(h, carpeta / f"{nombre}.pdf", perfil)
        print("listo", nombre)


if __name__ == "__main__":
    main()
